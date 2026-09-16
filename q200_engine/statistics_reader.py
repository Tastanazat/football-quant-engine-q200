"""
Q200 Engine - Statistics Reader

Q200 V3.1

CSV / XLSX dosyalarından statistics verisini okur.

Akış:

CSV / XLSX
    ↓
Statistics Reader
    ↓
Raw Statistics Dictionary
    ↓
Statistics Adapter
    ↓
Input Validation
    ↓
TeamStats

Bu katman:
- Model hesabı yapmaz.
- Odds kullanmaz.
- Model olasılığı üretmez.
- Statistics değerlerini değiştirmez.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from .input_validation import ALLOWED_FIELDS


# =========================================================
# SUPPORTED FILE TYPES
# =========================================================

SUPPORTED_EXTENSIONS = {
    ".csv",
    ".xlsx",
}


# =========================================================
# FIELD NORMALIZATION
# =========================================================

def _normalize_header(
    header: Any,
) -> str:

    if not isinstance(header, str):
        raise TypeError(
            "CSV/XLSX sütun isimleri string olmalıdır."
        )

    normalized = (
        header
        .strip()
        .lower()
    )

    if not normalized:
        raise ValueError(
            "Boş sütun adı kullanılamaz."
        )

    if normalized not in ALLOWED_FIELDS:
        raise ValueError(
            f"Bilinmeyen statistics sütunu: {normalized}"
        )

    return normalized


# =========================================================
# VALUE NORMALIZATION
# =========================================================

def _normalize_value(
    value: Any,
) -> Any:

    if isinstance(value, str):

        stripped = value.strip()

        if stripped == "":
            return None

        return stripped

    return value


# =========================================================
# CSV READER
# =========================================================

def read_csv(
    path: str | Path,
) -> list[dict[str, Any]]:

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Statistics dosyası bulunamadı: {file_path}"
        )

    if not file_path.is_file():
        raise ValueError(
            f"Statistics path bir dosya olmalıdır: {file_path}"
        )

    rows: list[dict[str, Any]] = []

    with file_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.DictReader(
            file
        )

        if reader.fieldnames is None:
            raise ValueError(
                "CSV dosyasında sütun başlığı bulunamadı."
            )

        headers = [
            _normalize_header(
                header
            )
            for header in reader.fieldnames
        ]

        if len(headers) != len(
            set(headers)
        ):
            raise ValueError(
                "CSV dosyasında tekrar eden sütun adı var."
            )

        for raw_row in reader:

            row: dict[str, Any] = {}

            for original_header, value in raw_row.items():

                if original_header is None:
                    continue

                header = _normalize_header(
                    original_header
                )

                row[header] = _normalize_value(
                    value
                )

            # Tamamen boş satırları atla.
            if any(
                value is not None
                for value in row.values()
            ):
                rows.append(row)

    if not rows:
        raise ValueError(
            "CSV dosyasında veri satırı bulunamadı."
        )

    return rows


# =========================================================
# XLSX READER
# =========================================================

def read_xlsx(
    path: str | Path,
) -> list[dict[str, Any]]:

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Statistics dosyası bulunamadı: {file_path}"
        )

    if not file_path.is_file():
        raise ValueError(
            f"Statistics path bir dosya olmalıdır: {file_path}"
        )

    try:

        from openpyxl import load_workbook

    except ImportError as exc:

        raise ImportError(
            "XLSX okumak için openpyxl gereklidir."
        ) from exc

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    try:

        worksheet = workbook.active

        rows_iterator = worksheet.iter_rows(
            values_only=True
        )

        try:

            raw_headers = next(
                rows_iterator
            )

        except StopIteration as exc:

            raise ValueError(
                "XLSX dosyası boş."
            ) from exc

        headers = [
            _normalize_header(
                header
            )
            for header in raw_headers
            if header is not None
        ]

        if not headers:
            raise ValueError(
                "XLSX dosyasında sütun başlığı bulunamadı."
            )

        if len(headers) != len(
            set(headers)
        ):
            raise ValueError(
                "XLSX dosyasında tekrar eden sütun adı var."
            )

        rows: list[dict[str, Any]] = []

        for raw_row in rows_iterator:

            row: dict[str, Any] = {}

            for index, header in enumerate(
                headers
            ):

                value = (
                    raw_row[index]
                    if index < len(raw_row)
                    else None
                )

                row[header] = _normalize_value(
                    value
                )

            if any(
                value is not None
                for value in row.values()
            ):
                rows.append(row)

        if not rows:
            raise ValueError(
                "XLSX dosyasında veri satırı bulunamadı."
            )

        return rows

    finally:

        workbook.close()


# =========================================================
# GENERIC FILE READER
# =========================================================

def read_statistics_file(
    path: str | Path,
) -> list[dict[str, Any]]:

    file_path = Path(path)

    extension = (
        file_path.suffix
        .lower()
    )

    if extension == ".csv":

        return read_csv(
            file_path
        )

    if extension == ".xlsx":

        return read_xlsx(
            file_path
        )

    raise ValueError(
        "Desteklenmeyen statistics dosya tipi: "
        f"{extension or '[uzantı yok]'}"
    )


# =========================================================
# SINGLE RECORD
# =========================================================

def read_statistics_record(
    path: str | Path,
    row_index: int = 0,
) -> dict[str, Any]:

    if not isinstance(
        row_index,
        int,
    ):
        raise TypeError(
            "row_index integer olmalıdır."
        )

    if row_index < 0:
        raise ValueError(
            "row_index negatif olamaz."
        )

    rows = read_statistics_file(
        path
    )

    if row_index >= len(rows):
        raise IndexError(
            f"row_index {row_index} mevcut değil. "
            f"Toplam satır: {len(rows)}"
        )

    return dict(
        rows[row_index]
    )
