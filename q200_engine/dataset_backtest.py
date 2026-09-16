"""
Q200 Engine - Dataset Backtest Runner

Q200 V3.1

CSV/XLSX manifest dosyasındaki geçmiş maçları sırayla Q200 analizinden
geçirir, History'ye kaydeder, gerçek sonucu işler, settlement üretir ve
sonunda Performance + Calibration + Evaluation raporu oluşturur.

Post-match sonuçlar yalnızca analiz tamamlandıktan sonra History/Settlement
katmanına aktarılır; model hesabına geri beslenmez.
"""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from .analyzer import run_q200_from_files
from .evaluation import EvaluationReport, build_evaluation
from .history import AnalysisHistory


DATASET_BACKTEST_VERSION = "Q200-DATASET-BACKTEST-V1"

_REQUIRED_HEADERS = {
    "match_id",
    "statistics_file_1",
    "statistics_file_2",
    "odds_file",
    "bankroll",
    "home_goals",
    "away_goals",
}

_OPTIONAL_HEADERS = {
    "uncertainty",
    "row_index_1",
    "row_index_2",
}

_ALLOWED_HEADERS = (
    _REQUIRED_HEADERS
    | _OPTIONAL_HEADERS
)


@dataclass(frozen=True)
class DatasetBacktestItem:
    """Tek bir dataset satırının sonucu."""

    row_number: int
    match_id: str
    record_id: int
    home_goals: int
    away_goals: int
    settlement_recorded: bool
    total_bets: int
    wins: int
    losses: int
    voids: int
    total_stake: float
    total_profit: float


@dataclass(frozen=True)
class DatasetBacktestSummary:
    """Toplu dataset backtest sonucu."""

    requested_rows: int
    processed_rows: int
    failed_rows: int
    results: tuple[DatasetBacktestItem, ...]
    evaluation: EvaluationReport


def _normalize_header(value: Any) -> str:

    if not isinstance(value, str):
        raise TypeError(
            "Dataset sütun isimleri string olmalıdır."
        )

    header = value.strip().lower()

    if not header:
        raise ValueError(
            "Dataset sütun adı boş olamaz."
        )

    if header not in _ALLOWED_HEADERS:
        raise ValueError(
            f"Bilinmeyen dataset sütunu: {header}"
        )

    return header


def _normalize_row(
    raw_row: dict[str, Any],
) -> dict[str, Any]:

    row: dict[str, Any] = {}

    for key, value in raw_row.items():

        if key is None:
            continue

        row[
            _normalize_header(key)
        ] = value

    missing = sorted(
        _REQUIRED_HEADERS
        - set(row)
    )

    if missing:
        raise ValueError(
            "Dataset satırında eksik alan: "
            + ", ".join(missing)
        )

    return row


def _read_csv(
    path: Path,
) -> list[dict[str, Any]]:

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.DictReader(
            file
        )

        if reader.fieldnames is None:
            raise ValueError(
                "Dataset CSV sütun başlığı içermiyor."
            )

        headers = [
            _normalize_header(
                header
            )
            for header in reader.fieldnames
        ]

        if len(headers) != len(
            set(headers)
        ):
            raise ValueError(
                "Dataset dosyasında tekrar eden sütun adı var."
            )

        rows = []

        for raw_row in reader:

            row = _normalize_row(
                raw_row
            )

            if any(
                value is not None
                and str(value).strip() != ""
                for value in row.values()
            ):
                rows.append(row)

        if not rows:
            raise ValueError(
                "Dataset dosyasında veri satırı bulunamadı."
            )

        return rows


def _read_xlsx(
    path: Path,
) -> list[dict[str, Any]]:

    try:

        from openpyxl import (
            load_workbook
        )

    except ImportError as exc:

        raise ImportError(
            "XLSX dataset okumak için openpyxl gereklidir."
        ) from exc

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:

        worksheet = workbook.active

        iterator = worksheet.iter_rows(
            values_only=True
        )

        try:

            raw_headers = next(
                iterator
            )

        except StopIteration as exc:

            raise ValueError(
                "Dataset XLSX boş."
            ) from exc

        headers = [
            _normalize_header(
                header
            )
            for header in raw_headers
            if header is not None
        ]

        if not headers:
            raise ValueError(
                "Dataset XLSX sütun başlığı içermiyor."
            )

        if len(headers) != len(
            set(headers)
        ):
            raise ValueError(
                "Dataset dosyasında tekrar eden sütun adı var."
            )

        rows = []

        for raw_row in iterator:

            row = {}

            for index, header in enumerate(
                headers
            ):

                value = (
                    raw_row[index]
                    if index < len(raw_row)
                    else None
                )

                row[
                    header
                ] = value

            if any(
                value is not None
                and str(value).strip() != ""
                for value in row.values()
            ):
                rows.append(
                    _normalize_row(
                        row
                    )
                )

        if not rows:
            raise ValueError(
                "Dataset XLSX veri satırı içermiyor."
            )

        return rows

    finally:

        workbook.close()


def read_dataset_file(
    path: str | Path,
) -> list[dict[str, Any]]:
    """CSV/XLSX dataset manifestini okur."""

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Dataset dosyası bulunamadı: {file_path}"
        )

    if not file_path.is_file():
        raise ValueError(
            f"Dataset path bir dosya olmalıdır: {file_path}"
        )

    extension = file_path.suffix.lower()

    if extension == ".csv":

        return _read_csv(
            file_path
        )

    if extension == ".xlsx":

        return _read_xlsx(
            file_path
        )

    raise ValueError(
        "Desteklenmeyen dataset dosya tipi: "
        f"{extension or '[uzantı yok]'}"
    )


def _text(
    value: Any,
    name: str,
) -> str:

    if value is None:
        raise ValueError(
            f"{name} boş olamaz."
        )

    result = str(value).strip()

    if not result:
        raise ValueError(
            f"{name} boş olamaz."
        )

    return result


def _integer(
    value: Any,
    name: str,
) -> int:

    if isinstance(
        value,
        bool,
    ):
        raise TypeError(
            f"{name} integer olmalıdır."
        )

    try:

        result = int(
            value
        )

    except (
        TypeError,
        ValueError,
    ) as exc:

        raise ValueError(
            f"{name} integer olmalıdır."
        ) from exc

    if not isinstance(
        value,
        int,
    ):

        try:

            if float(value) != result:
                raise ValueError

        except (
            TypeError,
            ValueError,
        ) as exc:

            raise ValueError(
                f"{name} integer olmalıdır."
            ) from exc

    if result < 0:
        raise ValueError(
            f"{name} negatif olamaz."
        )

    return result


def _positive_float(
    value: Any,
    name: str,
) -> float:

    if isinstance(
        value,
        bool,
    ):
        raise TypeError(
            f"{name} sayısal olmalıdır."
        )

    try:

        result = float(
            value
        )

    except (
        TypeError,
        ValueError,
    ) as exc:

        raise ValueError(
            f"{name} sayısal olmalıdır."
        ) from exc

    if not math.isfinite(
        result
    ):
        raise ValueError(
            f"{name} sonlu bir sayı olmalıdır."
        )

    if result <= 0:
        raise ValueError(
            f"{name} pozitif olmalıdır."
        )

    return result


def _optional_integer(
    value: Any,
    default: int,
    name: str,
) -> int:

    if (
        value is None
        or str(value).strip() == ""
    ):
        return default

    return _integer(
        value,
        name,
    )


class DatasetBacktestRunner:
    """
    Dataset manifestini:

        Analysis
        →
        History
        →
        Result
        →
        Settlement
        →
        Evaluation

    zincirinden geçirir.
    """

    def __init__(
        self,
        history_path: str | Path,
    ) -> None:

        self.history = AnalysisHistory(
            history_path
        )

    @staticmethod
    def _resolve_path(
        value: Any,
        dataset_path: Path,
        name: str,
    ) -> Path:

        raw = Path(
            _text(
                value,
                name,
            )
        )

        if not raw.is_absolute():

            raw = (
                dataset_path.parent
                / raw
            )

        return raw

    def run_rows(
        self,
        rows: Iterable[
            dict[str, Any]
        ],
        *,
        dataset_path: str | Path,
        continue_on_error: bool = False,
        bucket_count: int = 10,
    ) -> DatasetBacktestSummary:
        """Hazır dataset satırlarını çalıştırır."""

        if isinstance(
            rows,
            (
                str,
                bytes,
                dict,
            ),
        ):
            raise TypeError(
                "rows dictionary kayıtlarından oluşan iterable "
                "olmalıdır."
            )

        if not isinstance(
            continue_on_error,
            bool,
        ):
            raise TypeError(
                "continue_on_error boolean olmalıdır."
            )

        dataset_file = Path(
            dataset_path
        )

        items = list(
            rows
        )

        results: list[
            DatasetBacktestItem
        ] = []

        failed = 0

        for row_number, raw_row in enumerate(
            items,
            start=2,
        ):

            try:

                row = _normalize_row(
                    raw_row
                )

                match_id = _text(
                    row["match_id"],
                    "match_id",
                )

                statistics_file_1 = (
                    self._resolve_path(
                        row[
                            "statistics_file_1"
                        ],
                        dataset_file,
                        "statistics_file_1",
                    )
                )

                statistics_file_2 = (
                    self._resolve_path(
                        row[
                            "statistics_file_2"
                        ],
                        dataset_file,
                        "statistics_file_2",
                    )
                )

                odds_file = (
                    self._resolve_path(
                        row["odds_file"],
                        dataset_file,
                        "odds_file",
                    )
                )

                bankroll = _positive_float(
                    row["bankroll"],
                    "bankroll",
                )

                uncertainty = _text(
                    row.get(
                        "uncertainty",
                        "MEDIUM",
                    )
                    or "MEDIUM",
                    "uncertainty",
                ).upper()

                row_index_1 = (
                    _optional_integer(
                        row.get(
                            "row_index_1"
                        ),
                        0,
                        "row_index_1",
                    )
                )

                row_index_2 = (
                    _optional_integer(
                        row.get(
                            "row_index_2"
                        ),
                        0,
                        "row_index_2",
                    )
                )

                home_goals = _integer(
                    row["home_goals"],
                    "home_goals",
                )

                away_goals = _integer(
                    row["away_goals"],
                    "away_goals",
                )

                # =================================================
                # ANALYSIS
                # =================================================

                analysis = run_q200_from_files(
                    statistics_file_1=(
                        statistics_file_1
                    ),
                    statistics_file_2=(
                        statistics_file_2
                    ),
                    odds_file=odds_file,
                    bankroll=bankroll,
                    uncertainty=uncertainty,
                    row_index_1=row_index_1,
                    row_index_2=row_index_2,
                )

                # =================================================
                # HISTORY
                # =================================================

                record_id = self.history.save(
                    analysis,
                    match_id,
                )

                # =================================================
                # RESULT
                # =================================================

                recorded = (
                    self.history.record_result(
                        record_id,
                        home_goals,
                        away_goals,
                    )
                )

                if not recorded:
                    raise ValueError(
                        "History sonucu kaydedilemedi: "
                        f"{record_id}"
                    )

                # =================================================
                # SETTLEMENT
                # =================================================

                settlement = (
                    self.history.settle_record(
                        record_id
                    )
                )

                summary = settlement.get(
                    "summary"
                )

                if not isinstance(
                    summary,
                    dict,
                ):
                    raise ValueError(
                        "Settlement summary dictionary olmalıdır."
                    )

                results.append(
                    DatasetBacktestItem(
                        row_number=row_number,
                        match_id=match_id,
                        record_id=record_id,
                        home_goals=home_goals,
                        away_goals=away_goals,
                        settlement_recorded=True,
                        total_bets=int(
                            summary.get(
                                "total_bets",
                                0,
                            )
                        ),
                        wins=int(
                            summary.get(
                                "wins",
                                0,
                            )
                        ),
                        losses=int(
                            summary.get(
                                "losses",
                                0,
                            )
                        ),
                        voids=int(
                            summary.get(
                                "voids",
                                0,
                            )
                        ),
                        total_stake=float(
                            summary.get(
                                "total_stake",
                                0.0,
                            )
                        ),
                        total_profit=float(
                            summary.get(
                                "total_profit",
                                0.0,
                            )
                        ),
                    )
                )

            except (
                TypeError,
                ValueError,
                FileNotFoundError,
                IndexError,
                ImportError,
            ):

                failed += 1

                if not continue_on_error:
                    raise

        # =========================================================
        # FINAL EVALUATION
        # =========================================================

        evaluation = build_evaluation(
            self.history,
            bucket_count=bucket_count,
        )

        return DatasetBacktestSummary(
            requested_rows=len(
                items
            ),
            processed_rows=len(
                results
            ),
            failed_rows=failed,
            results=tuple(
                results
            ),
            evaluation=evaluation,
        )

    def run_file(
        self,
        dataset_path: str | Path,
        *,
        continue_on_error: bool = False,
        bucket_count: int = 10,
    ) -> DatasetBacktestSummary:
        """CSV/XLSX manifestini okuyup çalıştırır."""

        path = Path(
            dataset_path
        )

        rows = read_dataset_file(
            path
        )

        return self.run_rows(
            rows,
            dataset_path=path,
            continue_on_error=(
                continue_on_error
            ),
            bucket_count=bucket_count,
        )


def run_dataset_backtest(
    dataset_path: str | Path,
    history_path: str | Path,
    *,
    continue_on_error: bool = False,
    bucket_count: int = 10,
) -> DatasetBacktestSummary:
    """Dataset Backtest Runner için kısa kullanım wrapper'ı."""

    return DatasetBacktestRunner(
        history_path
    ).run_file(
        dataset_path,
        continue_on_error=(
            continue_on_error
        ),
        bucket_count=bucket_count,
    )
